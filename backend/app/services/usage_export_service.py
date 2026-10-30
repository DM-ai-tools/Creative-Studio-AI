"""Build Excel usage reports from usage_events."""

from __future__ import annotations

import io
from datetime import date, datetime, time, timedelta, timezone
from typing import Any
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from app.core.config import settings
from app.core.database import AsyncSessionLocal
from app.models.asset import Asset
from app.models.brand import Brand
from app.models.brief import Brief
from app.models.tenant import Tenant
from app.models.user import User
from app.models.variant import Variant
from app.services.public_urls import (
    absolute_media_url,
    media_urls_from_extra,
    normalize_stored_file_url,
    pick_image_and_video_links,
    public_app_origin,
    server_disk_path,
    variant_library_url,
)
from app.services.usage_tracker import (
    _display_model,
    effective_cost_usd,
    fetch_usage_events,
    is_byteplus_seedance_row,
)


MAX_EXPORT_ROWS = 25_000
MAX_RANGE_DAYS = 366

DETAIL_HEADERS = [
    "Date (UTC)",
    "User email",
    "User name",
    "Client / workspace",
    "Brand ID",
    "Brand name",
    "Category",
    "Media type",
    "Image link (for manager — click to open)",
    "Video link (for manager — click to open)",
    "File URL in app (/files/...)",
    "Server file location (local dev only)",
    "Creative library (open in app)",
    "Provider",
    "Model",
    "Operation",
    "Duration (sec)",
    "Prompt tokens",
    "Completion tokens",
    "Total tokens",
    "Cost (USD)",
    "Cost estimated",
    "Status",
    "Error",
    "Brief ID",
    "Variant ID",
    "Job ID",
    "Task ID",
    "Event ID",
]

# Column indices for summary rollups (0-based, must match DETAIL_HEADERS order)
_COL_STATUS = DETAIL_HEADERS.index("Status")
_COL_TOTAL_TOKENS = DETAIL_HEADERS.index("Total tokens")
_COL_COST = DETAIL_HEADERS.index("Cost (USD)")
_COL_MODEL = DETAIL_HEADERS.index("Model")
_COL_USER_EMAIL = DETAIL_HEADERS.index("User email")
_COL_BRAND_NAME = DETAIL_HEADERS.index("Brand name")
_COL_HYPERLINK_HEADERS = (
    "Image link (for manager — click to open)",
    "Video link (for manager — click to open)",
    "Creative library (open in app)",
)


def _uuid_or_none(value: Any) -> UUID | None:
    try:
        return UUID(str(value)) if value else None
    except (TypeError, ValueError, AttributeError):
        return None


def _load_creative_studio_job_metadata(
    job_ids: set[str],
    *,
    job_tenants: dict[str, str] | None = None,
) -> dict[str, dict[str, str]]:
    """Recover report fields from persisted Creative Studio jobs.

    Usage events created by older builds sometimes contain only ``job_id``.  The
    persisted job is the authoritative source for the selected user/brand and
    for the final stitched output URL, so keep those values together here.
    """
    if not job_ids:
        return {}
    from app.services.creative_studio_job_service import _load_job_from_disk

    out: dict[str, dict[str, str]] = {}
    for jid in job_ids:
        if not jid:
            continue
        # Keep tenant exports isolated even though persisted jobs live below a
        # shared uploads root.
        job = _load_job_from_disk(jid, tenant_id=(job_tenants or {}).get(jid))
        if not job:
            continue
        payload = job.get("payload") if isinstance(job.get("payload"), dict) else {}
        result = job.get("result") if isinstance(job.get("result"), dict) else {}
        out[jid] = {
            "url": str(result.get("url") or "").strip(),
            "user_id": str(payload.get("user_id") or "").strip(),
            "brand_id": str(payload.get("brand_id") or "").strip(),
            "brand_name": str(payload.get("brand_name") or "").strip(),
            "brief_id": str(payload.get("brief_id") or "").strip(),
            "media_mode": str(payload.get("media_mode") or result.get("media_mode") or "").strip(),
        }
    return out


def _job_urls(job_metadata: dict[str, dict[str, str]]) -> dict[str, str]:
    return {
        jid: str(meta.get("url") or "").strip()
        for jid, meta in job_metadata.items()
        if str(meta.get("url") or "").strip()
    }


def _stored_paths_from_extra(extra: dict | None, *, job_urls: dict[str, str]) -> list[str]:
    paths: list[str] = []
    for abs_u in media_urls_from_extra(extra):
        norm = normalize_stored_file_url(abs_u)
        if norm:
            paths.append(norm)
    jid = str((extra or {}).get("job_id") or "").strip()
    if jid and jid in job_urls:
        norm = normalize_stored_file_url(job_urls[jid])
        if norm:
            paths.append(norm)
    return paths


def _category(operation: str, provider: str) -> str:
    op = (operation or "").lower()
    prov = (provider or "").lower()
    if "creative_studio" in op or prov == "byteplus":
        return "Creative Studio"
    if op in {
        "generate_image",
        "hero_ai_image",
        "analyze_reference_image",
        "creative_studio_image",
        "creative_studio_character_portrait",
    } or "image" in op:
        return "Image generation"
    if op == "generate_video" or "video" in op:
        return "Video generation"
    if prov == "firecrawl" or "fetch_brand" in op:
        return "Brand scrape"
    if prov == "openrouter" or op.startswith("llm") or "chat" in op:
        return "LLM / copy"
    if prov == "heygen":
        return "HeyGen"
    if prov == "higgsfield":
        return "Higgsfield"
    if prov == "sociavault":
        return "Social scrape"
    if prov == "meta":
        return "Meta export"
    return "Other"


def _media_type(operation: str, provider: str, extra: dict | None) -> str:
    mt = str((extra or {}).get("media_type") or "").strip().lower()
    if mt in {"image", "video", "text"}:
        return mt.title()
    op = (operation or "").lower()
    if "video" in op or (extra or {}).get("video_url"):
        return "Video"
    if "image" in op or (extra or {}).get("image_url"):
        return "Image"
    if prov := (provider or "").lower():
        if prov == "byteplus" and "video" in op:
            return "Video"
    if op.startswith("llm") or "chat" in op or prov == "openrouter":
        return "Text"
    return ""


def parse_usage_date_range(
    *,
    days: int | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
) -> tuple[datetime, datetime, str]:
    now = datetime.now(timezone.utc)
    if date_from and date_to:
        if date_to < date_from:
            raise ValueError("End date must be on or after start date")
        span = (date_to - date_from).days + 1
        if span > MAX_RANGE_DAYS:
            raise ValueError(f"Date range cannot exceed {MAX_RANGE_DAYS} days")
        since = datetime.combine(date_from, time.min, tzinfo=timezone.utc)
        until = datetime.combine(date_to, time.max, tzinfo=timezone.utc)
        label = f"{date_from.isoformat()}_{date_to.isoformat()}"
        return since, until, label
    period = max(1, min(int(days or 30), MAX_RANGE_DAYS))
    since = now - timedelta(days=period)
    until = now
    label = f"last_{period}_days"
    return since, until, label


def _detail_row(
    row: Any,
    *,
    users: dict[UUID, User],
    tenants: dict[UUID, str],
    brands: dict[UUID, str],
    brief_brands: dict[UUID, UUID] | None = None,
    variant_brands: dict[UUID, UUID] | None = None,
    variant_asset_urls: dict[UUID, list[str]] | None = None,
    job_metadata: dict[str, dict[str, str]] | None = None,
    file_url_to_variant: dict[str, UUID] | None = None,
    file_url_to_brand: dict[str, UUID] | None = None,
    file_url_to_user: dict[str, UUID] | None = None,
    pipeline_file_to_variant: dict[str, UUID] | None = None,
    variant_pipeline_media: dict[UUID, dict[str, str]] | None = None,
) -> list[Any]:
    extra = row.extra or {}
    job_id_str = str(extra.get("job_id") or "").strip()
    job_meta = (job_metadata or {}).get(job_id_str) or {}
    job_url = str(job_meta.get("url") or "").strip()
    job_media_urls = {job_id_str: job_url} if job_id_str and job_url else {}
    brief_id_str = str(extra.get("brief_id") or job_meta.get("brief_id") or "").strip()
    brand_id_str = str(extra.get("brand_id") or job_meta.get("brand_id") or "").strip()
    brand_name = str(extra.get("brand_name") or job_meta.get("brand_name") or "").strip()

    cost = effective_cost_usd(row)
    estimated = cost > 0 and float(row.cost_usd or 0) <= 0 and is_byteplus_seedance_row(row)
    urls = media_urls_from_extra(extra)
    # A provider row can point at an intermediate segment.  Put the persisted
    # Creative Studio result first so a manager opens the final stitched asset.
    final_job_url = absolute_media_url(job_url)
    if final_job_url:
        urls = [final_job_url, *[url for url in urls if url != final_job_url]]
    variant_id_str = str(extra.get("variant_id") or "").strip()
    stored_paths = _stored_paths_from_extra(extra, job_urls=job_media_urls or {})
    stored_file = stored_paths[0] if stored_paths else ""
    if not variant_id_str and file_url_to_variant:
        for path in stored_paths:
            vid = file_url_to_variant.get(path)
            if vid:
                variant_id_str = str(vid)
                break
    if not variant_id_str and stored_file and pipeline_file_to_variant:
        vid = pipeline_file_to_variant.get(stored_file)
        if vid:
            variant_id_str = str(vid)

    # Resolve the brand only after asset/pipeline lookups have had a chance to
    # recover a missing variant.  This order is important for historical rows.
    brand_uuid = _uuid_or_none(brand_id_str)
    if not brand_uuid and brief_id_str and brief_brands:
        brand_uuid = brief_brands.get(_uuid_or_none(brief_id_str))
    if not brand_uuid and variant_id_str and variant_brands:
        brand_uuid = variant_brands.get(_uuid_or_none(variant_id_str))
    if not brand_uuid and file_url_to_brand:
        for path in stored_paths:
            if file_url_to_brand.get(path):
                brand_uuid = file_url_to_brand[path]
                break
    if brand_uuid:
        brand_id_str = str(brand_uuid)
        if not brand_name:
            brand_name = brands.get(brand_uuid, "")

    user_uuid = row.user_id or _uuid_or_none(job_meta.get("user_id"))
    if not user_uuid and file_url_to_user:
        for path in stored_paths:
            if file_url_to_user.get(path):
                user_uuid = file_url_to_user[path]
                break
    user = users.get(user_uuid) if user_uuid else None
    if not urls and variant_asset_urls and variant_id_str:
        try:
            urls = list(variant_asset_urls.get(UUID(variant_id_str)) or [])
        except ValueError:
            pass
    if not stored_file and urls:
        stored_file = normalize_stored_file_url(urls[0])
    if not urls and stored_file:
        urls = [absolute_media_url(stored_file)]
    image_link, video_link = pick_image_and_video_links(
        urls,
        extra=extra,
        operation=row.operation or "",
        provider=row.provider or "",
    )
    if variant_id_str and variant_pipeline_media:
        try:
            pipe = variant_pipeline_media.get(UUID(variant_id_str)) or {}
            if not image_link and pipe.get("image"):
                image_link = pipe["image"]
            if not video_link and pipe.get("video"):
                video_link = pipe["video"]
        except ValueError:
            pass
    media_type = _media_type(row.operation, row.provider, extra)
    if not brand_name and media_type in {"Image", "Video"}:
        brand_name = "Legacy – brand unavailable"
    disk_path = ""
    if not settings.is_production and stored_file:
        disk_path = server_disk_path(stored_file)
    library_link = variant_library_url(variant_id_str) if variant_id_str else ""
    prompt_tok = int(row.prompt_tokens or 0)
    completion_tok = int(row.completion_tokens or 0)
    total_tok = int(row.total_tokens or 0)
    if total_tok and not completion_tok and is_byteplus_seedance_row(row):
        completion_tok = total_tok if not prompt_tok else max(0, total_tok - prompt_tok)
    task_id = extra.get("task_id") if isinstance(extra.get("task_id"), str) else ""
    if not task_id and isinstance(extra, dict):
        task_id = str(extra.get("task_id") or "")

    duration = extra.get("duration_seconds")
    try:
        duration_val = float(duration) if duration is not None else ""
    except (TypeError, ValueError):
        duration_val = ""

    created = row.created_at.astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M:%S") if row.created_at else ""

    return [
        created,
        (user.email if user else ""),
        (user.full_name if user else ""),
        tenants.get(row.tenant_id, "") if row.tenant_id else "",
        brand_id_str,
        brand_name,
        _category(row.operation, row.provider),
        media_type,
        image_link,
        video_link,
        stored_file,
        disk_path,
        library_link,
        row.provider,
        _display_model(row.model, row.provider),
        row.operation,
        duration_val,
        prompt_tok,
        completion_tok,
        total_tok or (prompt_tok + completion_tok),
        round(cost, 6),
        "Yes" if estimated else "No",
        "Success" if row.success else "Failed",
        (row.error or "")[:500] if row.error else "",
        brief_id_str,
        variant_id_str,
        job_id_str,
        task_id,
        str(row.id),
    ]


def _write_sheet_headers(ws, headers: list[str]) -> None:
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True)


def _autosize_columns(ws, max_width: int = 48) -> None:
    for col_idx, column_cells in enumerate(ws.columns, 1):
        length = max(len(str(c.value or "")) for c in column_cells)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(length + 2, 10), max_width)


def _add_hyperlinks(ws, url_col_indices: list[int], start_row: int = 2) -> None:
    link_font = Font(color="0563C1", underline="single")
    for row_idx in range(start_row, ws.max_row + 1):
        for col_idx in url_col_indices:
            cell = ws.cell(row=row_idx, column=col_idx)
            val = str(cell.value or "").strip()
            if val.startswith("http"):
                cell.hyperlink = val
                cell.font = link_font


def _summary_rows(
    detail_rows: list[list[Any]],
    *,
    group_col: int,
    label: str,
) -> list[list[Any]]:
    acc: dict[str, dict[str, Any]] = {}
    for row in detail_rows:
        key = str(row[group_col] or "(blank)")
        slot = acc.setdefault(key, {"calls": 0, "failed": 0, "cost": 0.0, "tokens": 0})
        slot["calls"] += 1
        if row[_COL_STATUS] == "Failed":
            slot["failed"] += 1
        try:
            slot["cost"] += float(row[_COL_COST] or 0)
        except (TypeError, ValueError):
            pass
        slot["tokens"] += int(row[_COL_TOTAL_TOKENS] or 0)
    out = [[label, "Calls", "Failed", "Total tokens", "Cost (USD)"]]
    for key, slot in sorted(acc.items(), key=lambda x: x[1]["cost"], reverse=True):
        out.append([key, slot["calls"], slot["failed"], slot["tokens"], round(slot["cost"], 4)])
    return out


def _readme_sheet_rows() -> list[list[str]]:
    origin = public_app_origin()
    prod = settings.is_production
    return [
        ["Usage report — for your reporting manager"],
        [""],
        ["How to use", "Log in as admin on the live app → Usage & costs → Download Excel → email this file."],
        ["Image link", f"Blue hyperlink — opens the still in a browser ({origin}/files/...)."],
        ["Video link", f"Blue hyperlink — opens the video in a browser ({origin}/files/...)."],
        ["Creative library", "Optional — opens that saved creative inside the app (variants page)."],
        [""],
        ["Production (creativestudio.tools)", "Send Image link + Video link columns only; disk path is left blank on purpose."],
        ["Local development", "Disk path column shows where the file sits under backend/uploads on your PC."],
        [""],
        ["Empty image/video cell?", "That row was text/LLM only, or the file was not recorded yet. Re-export after new work."],
        ["Tip", "Save Creative Studio outputs to Creative library so variant links fill in for managers."],
    ]


def build_usage_workbook(detail_rows: list[list[Any]]) -> bytes:
    wb = Workbook()
    readme = wb.active
    readme.title = "Read me"
    for row in _readme_sheet_rows():
        readme.append(row)
    readme.column_dimensions["A"].width = 36
    readme.column_dimensions["B"].width = 72

    ws = wb.create_sheet(title="Usage detail", index=1)
    _write_sheet_headers(ws, DETAIL_HEADERS)
    for row in detail_rows:
        ws.append(row)
    link_cols = [DETAIL_HEADERS.index(h) + 1 for h in _COL_HYPERLINK_HEADERS]
    _add_hyperlinks(ws, url_col_indices=link_cols)
    _autosize_columns(ws)

    for title, group_col in (
        ("By user", _COL_USER_EMAIL),
        ("By brand", _COL_BRAND_NAME),
        ("By model", _COL_MODEL),
    ):
        sheet = wb.create_sheet(title=title)
        for row in _summary_rows(detail_rows, group_col=group_col, label=DETAIL_HEADERS[group_col]):
            sheet.append(row)
        _autosize_columns(sheet, max_width=36)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def generate_usage_export_xlsx(
    *,
    tenant_id: UUID | None,
    platform: bool,
    days: int | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    user_id: UUID | None = None,
) -> tuple[bytes, str]:
    since, until, label = parse_usage_date_range(
        days=days, date_from=date_from, date_to=date_to
    )
    rows = await fetch_usage_events(
        tenant_id=tenant_id,
        platform=platform,
        since=since,
        until=until,
        user_id=user_id,
        limit=MAX_EXPORT_ROWS + 1,
    )
    if len(rows) > MAX_EXPORT_ROWS:
        raise ValueError(
            f"Too many rows ({len(rows)}). Narrow the date range (max {MAX_EXPORT_ROWS} rows per export)."
        )

    user_ids = {r.user_id for r in rows if r.user_id}
    tenant_ids = {r.tenant_id for r in rows if r.tenant_id}
    brand_ids: set[UUID] = set()
    brief_ids: set[UUID] = set()
    variant_ids: set[UUID] = set()
    for r in rows:
        ex = r.extra or {}
        bid = ex.get("brand_id")
        if bid:
            try:
                brand_ids.add(UUID(str(bid)))
            except ValueError:
                pass
        for key, bucket in (("brief_id", brief_ids), ("variant_id", variant_ids)):
            raw = ex.get(key)
            if raw:
                try:
                    bucket.add(UUID(str(raw)))
                except ValueError:
                    pass

    users: dict[UUID, User] = {}
    tenants: dict[UUID, str] = {}
    brands: dict[UUID, str] = {}
    brief_brands: dict[UUID, UUID] = {}
    variant_brands: dict[UUID, UUID] = {}
    variant_asset_urls: dict[UUID, list[str]] = {}
    job_ids: set[str] = set()
    job_tenants: dict[str, str] = {}
    stored_file_paths: set[str] = set()
    for r in rows:
        ex = r.extra or {}
        jid = str(ex.get("job_id") or "").strip()
        if jid:
            job_ids.add(jid)
            if r.tenant_id:
                job_tenants[jid] = str(r.tenant_id)
        for path in _stored_paths_from_extra(ex, job_urls={}):
            stored_file_paths.add(path)

    job_metadata = _load_creative_studio_job_metadata(job_ids, job_tenants=job_tenants)
    job_media_urls = _job_urls(job_metadata)
    for meta in job_metadata.values():
        uid = _uuid_or_none(meta.get("user_id"))
        bid = _uuid_or_none(meta.get("brand_id"))
        brief_id = _uuid_or_none(meta.get("brief_id"))
        if uid:
            user_ids.add(uid)
        if bid:
            brand_ids.add(bid)
        if brief_id:
            brief_ids.add(brief_id)
    for raw in job_media_urls.values():
        norm = normalize_stored_file_url(raw)
        if norm:
            stored_file_paths.add(norm)

    file_url_to_variant: dict[str, UUID] = {}
    file_url_to_brand: dict[str, UUID] = {}
    file_url_to_user: dict[str, UUID] = {}
    pipeline_file_to_variant: dict[str, UUID] = {}
    variant_pipeline_media: dict[UUID, dict[str, str]] = {}

    async with AsyncSessionLocal() as db:
        if tenant_ids:
            trows = await db.execute(select(Tenant.id, Tenant.name).where(Tenant.id.in_(tenant_ids)))
            tenants = {row.id: row.name for row in trows.all()}
        if stored_file_paths:
            path_list = list(stored_file_paths)
            arows = await db.execute(
                select(Asset.file_url, Asset.variant_id, Asset.brand_id, Asset.created_by)
                .where(Asset.file_url.in_(path_list))
            )
            for file_url, variant_id, asset_brand_id, created_by in arows.all():
                if not file_url:
                    continue
                normalized = normalize_stored_file_url(str(file_url)) or str(file_url)
                if variant_id:
                    file_url_to_variant[normalized] = variant_id
                    variant_ids.add(variant_id)
                if asset_brand_id:
                    file_url_to_brand[normalized] = asset_brand_id
                    brand_ids.add(asset_brand_id)
                if created_by:
                    file_url_to_user[normalized] = created_by
                    user_ids.add(created_by)
        if tenant_ids:
            vpipe = await db.execute(
                select(Variant.id, Variant.brand_id, Variant.generation_params).where(
                    Variant.tenant_id.in_(tenant_ids)
                )
            )
            for vid, variant_brand_id, params in vpipe.all():
                if not isinstance(params, dict):
                    continue
                pipe = params.get("pipeline")
                if not isinstance(pipe, dict):
                    continue
                for block_key in ("image", "video"):
                    block = pipe.get(block_key)
                    if not isinstance(block, dict):
                        continue
                    norm = normalize_stored_file_url(str(block.get("url") or ""))
                    if norm:
                        pipeline_file_to_variant[norm] = vid
                        if norm in stored_file_paths:
                            variant_ids.add(vid)
                            if variant_brand_id:
                                variant_brands[vid] = variant_brand_id
                                brand_ids.add(variant_brand_id)
                img_u = absolute_media_url(str((pipe.get("image") or {}).get("url") or ""))
                vid_u = absolute_media_url(str((pipe.get("video") or {}).get("url") or ""))
                if img_u or vid_u:
                    variant_pipeline_media[vid] = {
                        "image": img_u,
                        "video": vid_u,
                    }
        if brief_ids:
            brows = await db.execute(select(Brief.id, Brief.brand_id).where(Brief.id.in_(brief_ids)))
            for row in brows.all():
                brief_brands[row.id] = row.brand_id
                brand_ids.add(row.brand_id)
        if variant_ids:
            vrows = await db.execute(select(Variant.id, Variant.brand_id).where(Variant.id.in_(variant_ids)))
            for row in vrows.all():
                variant_brands[row.id] = row.brand_id
                brand_ids.add(row.brand_id)
            arows = await db.execute(
                select(Asset.variant_id, Asset.file_url)
                .where(Asset.variant_id.in_(variant_ids), Asset.file_url.is_not(None))
                .order_by(Asset.created_at.desc())
            )
            for variant_id, file_url in arows.all():
                if not variant_id or not file_url:
                    continue
                slot = variant_asset_urls.setdefault(variant_id, [])
                abs_u = absolute_media_url(str(file_url))
                if abs_u and abs_u not in slot:
                    slot.append(abs_u)
        if user_ids:
            urows = await db.execute(select(User).where(User.id.in_(user_ids)))
            users = {u.id: u for u in urows.scalars().all()}
        if brand_ids:
            brows = await db.execute(select(Brand.id, Brand.name).where(Brand.id.in_(brand_ids)))
            brands = {row.id: row.name for row in brows.all()}

    detail = [
        _detail_row(
            r,
            users=users,
            tenants=tenants,
            brands=brands,
            brief_brands=brief_brands,
            variant_brands=variant_brands,
            variant_asset_urls=variant_asset_urls,
            job_metadata=job_metadata,
            file_url_to_variant=file_url_to_variant,
            file_url_to_brand=file_url_to_brand,
            file_url_to_user=file_url_to_user,
            pipeline_file_to_variant=pipeline_file_to_variant,
            variant_pipeline_media=variant_pipeline_media,
        )
        for r in reversed(rows)
    ]
    content = build_usage_workbook(detail)
    filename = f"usage-report-{label}.xlsx"
    return content, filename
