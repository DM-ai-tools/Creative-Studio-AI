"""Read client strategy uploads (.md, .txt, .docx) into markdown-like text."""

from __future__ import annotations

import io
import logging
import re

logger = logging.getLogger(__name__)

STRATEGY_FILE_EXTENSIONS = (".md", ".txt", ".markdown", ".docx", ".doc")


def strategy_filename_ok(filename: str) -> bool:
    name = (filename or "").lower()
    return any(name.endswith(ext) for ext in STRATEGY_FILE_EXTENSIONS)


def strategy_bytes_to_text(raw: bytes, *, filename: str = "") -> str:
    """Decode uploaded strategy bytes to plain / markdown text for parsing."""
    name = (filename or "strategy.md").lower()

    if name.endswith((".md", ".txt", ".markdown")):
        return _decode_text_bytes(raw)

    if name.endswith(".docx") or (name.endswith(".doc") and raw[:2] == b"PK"):
        return _docx_to_markdown(raw)

    if name.endswith(".doc"):
        raise ValueError(
            "Legacy .doc Word files are not supported. Save as .docx, .md, or .txt and upload again."
        )

    raise ValueError("Upload a strategy file (.md, .txt, .docx, or .doc)")


def _decode_text_bytes(raw: bytes) -> str:
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("latin-1", errors="replace")


def _docx_to_markdown(raw: bytes) -> str:
    text = ""
    try:
        import mammoth

        result = mammoth.convert_to_markdown(io.BytesIO(raw))
        text = (result.value or "").strip()
        if result.messages:
            for msg in result.messages:
                logger.info("docx conversion: %s", msg)
    except ImportError:
        logger.warning("mammoth not installed — using python-docx text extraction")

    if not text:
        text = _docx_plaintext_fallback(raw)

    text = text.strip()
    if not text:
        raise ValueError("Could not extract text from the Word document — the file may be empty.")
    return text


def _docx_plaintext_fallback(raw: bytes) -> str:
    """Fallback when mammoth is unavailable or returns empty (e.g. text boxes only)."""
    try:
        from docx import Document
    except ImportError as e:
        raise ValueError(
            "Word document support is not installed on the server. "
            "Run: pip install -r requirements.txt in the backend folder, then restart the backend."
        ) from e

    doc = Document(io.BytesIO(raw))
    parts: list[str] = []
    for para in doc.paragraphs:
        line = (para.text or "").strip()
        if line:
            parts.append(line)
    for table in doc.tables:
        for row in table.rows:
            cells = [re.sub(r"\s+", " ", (cell.text or "").strip()) for cell in row.cells]
            cells = [c for c in cells if c]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)
