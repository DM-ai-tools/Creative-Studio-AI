import io
from datetime import date, datetime, timezone
from types import SimpleNamespace
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from app.services.public_urls import absolute_media_url, media_urls_from_extra
from app.services.usage_export_service import (
    DETAIL_HEADERS,
    _detail_row,
    build_usage_workbook,
    parse_usage_date_range,
)


def test_parse_usage_date_range_days():
    since, until, label = parse_usage_date_range(days=30)
    assert until >= since
    assert label == "last_30_days"


def test_parse_usage_date_range_custom():
    since, until, label = parse_usage_date_range(
        date_from=date(2025, 1, 1),
        date_to=date(2025, 1, 31),
    )
    assert since.year == 2025
    assert until.day == 31
    assert "2025-01-01" in label


def test_parse_usage_date_range_rejects_inverted():
    with pytest.raises(ValueError):
        parse_usage_date_range(
            date_from=date(2025, 2, 1),
            date_to=date(2025, 1, 1),
        )


def test_extract_output_urls_dedupes():
    urls = media_urls_from_extra(
        {
            "image_url": "/files/a.png",
            "output_urls": ["/files/a.png", "https://x.test/b.mp4"],
        }
    )
    assert len(urls) == 2


def test_absolute_media_url():
    u = absolute_media_url("/files/test.png")
    assert u.endswith("/files/test.png")


def _usage_row(**overrides):
    values = {
        "id": uuid4(),
        "created_at": datetime(2026, 10, 29, 8, 30, tzinfo=timezone.utc),
        "user_id": None,
        "tenant_id": uuid4(),
        "extra": {},
        "provider": "byteplus",
        "model": "seedance-2.5",
        "operation": "video_generation",
        "prompt_tokens": 0,
        "completion_tokens": 100,
        "total_tokens": 100,
        "cost_usd": 1.25,
        "success": True,
        "error": None,
    }
    values.update(overrides)
    return SimpleNamespace(**values)


def test_detail_row_recovers_job_user_brand_and_final_video_link():
    user_id = uuid4()
    brand_id = uuid4()
    tenant_id = uuid4()
    row = _usage_row(tenant_id=tenant_id, extra={"job_id": "job-1"})
    detail = _detail_row(
        row,
        users={user_id: SimpleNamespace(email="creator@example.com", full_name="Creator")},
        tenants={tenant_id: "Client workspace"},
        brands={brand_id: "Traffic Radius"},
        job_metadata={
            "job-1": {
                "url": "/files/final-production.mp4",
                "user_id": str(user_id),
                "brand_id": str(brand_id),
                "brand_name": "Traffic Radius",
                "brief_id": "",
                "media_mode": "video",
            }
        },
    )

    assert detail[DETAIL_HEADERS.index("User email")] == "creator@example.com"
    assert detail[DETAIL_HEADERS.index("Brand name")] == "Traffic Radius"
    assert detail[DETAIL_HEADERS.index("Video link (for manager — click to open)")].endswith(
        "/files/final-production.mp4"
    )


def test_detail_row_recovers_brand_from_asset_path():
    brand_id = uuid4()
    row = _usage_row(
        provider="openai",
        model="gpt-image-2",
        operation="generate_image",
        extra={"image_url": "/files/legacy-image.png"},
    )
    detail = _detail_row(
        row,
        users={},
        tenants={},
        brands={brand_id: "Recovered brand"},
        file_url_to_brand={"/files/legacy-image.png": brand_id},
    )
    assert detail[DETAIL_HEADERS.index("Brand name")] == "Recovered brand"


def test_workbook_media_urls_are_clickable():
    row = [""] * len(DETAIL_HEADERS)
    video_url = "https://creativestudio.tools/files/report-video.mp4"
    row[DETAIL_HEADERS.index("Video link (for manager — click to open)")] = video_url
    content = build_usage_workbook([row])
    workbook = load_workbook(io.BytesIO(content))
    sheet = workbook["Usage detail"]
    col = DETAIL_HEADERS.index("Video link (for manager — click to open)") + 1
    assert sheet.cell(row=2, column=col).hyperlink.target == video_url
